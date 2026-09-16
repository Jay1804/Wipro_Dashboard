"""Export 'Client Wise All Cases - Bridge' from the Authbridge MIS Query Browser
and save the resulting CSV into the project root.
"""

import datetime
import io
import os
import zipfile

import pandas as pd
import requests
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

BASE_URL = "https://mis.authbridge.com/export_query"
ROOT_DIR = os.path.dirname(os.path.abspath(__file__))

FLEX_FIELDS_TO_DROP = [
    "case_flex_field27",
    "case_flex_field28",
    "case_flex_field29",
    "case_flex_field30",
]

NEW_HEADERS = [
    "company_name",
    "Candidate_name",
    "process_name",
    "location",
    "case_ars_no",
    "received_date",
    "case_created_date",
    "Resume ID*",
    "Employee Type",
    "Location Type (Ofshore/Onsite)",
    "BGV Type",
    "BGV Lead",
    "Value",
    "I-verify Type",
    "Value",
    "Entity Name",
    "Scope of Verification",
    "Account Name",
    "Country",
    "BGV Pattern Name",
    "BGV Assigned Date",
    "BGV1 Assigned Date",
    "BGV2 Assigned Date",
    "Scope of Verification 1",
    "Candidate Name As per I-Verify",
    "WIPRO DOJ",
    "IVERIFY SPOC NAME ",
    "Vault Search ID ",
    "I verify date & time",
    "GEO",
    "masterId",
    "API_Created_stamp",
    "Recruiter_notification_stamp",
    "case_due_date",
    "last_insuff_date",
    "last_insuff_fulfill_date",
    "Case_status",
    "DATA_SOURCE",
    "CASE_PRIORITY_FLAG",
    "DQC_RELEASED_DATE",
]

LOCATION_MAP = {
    "india": "India",
    "ind": "India",
    "inda": "India",
    "na": "India",
    "my": "Malaysia",
    "malaysia": "Malaysia",
    "japan": "Japan",
    "saudi arabia": "Saudi Arabia",
    "uae": "UAE",
}


def normalize_location(account_name):
    if not isinstance(account_name, str):
        return "India"
    cleaned = account_name.strip().strip('"').strip()
    if not cleaned:
        return "India"
    return LOCATION_MAP.get(cleaned.lower(), account_name)


USERNAME = os.environ["MIS_USERNAME"]
PASSWORD = os.environ["MIS_PASSWORD"]

LOGIN_PAYLOAD = {
    "username": USERNAME,
    "password": PASSWORD,
    "login": "Login",
}

EXPORT_PAYLOAD = {
    "hostname": "4",  # Bridge Live
    "database": "checkpoint_live",  # Bridge Live
    "access_time": "Tracker - Query",
    "csv_query": "263",  # Client Wise All Cases - Bridge
    "query_days_range": "",
    "client1": "3017",
}


def export_csv():
    session = requests.Session()
    session.headers.update({"User-Agent": "Mozilla/5.0"})

    login_resp = session.post(f"{BASE_URL}/login.php", data=LOGIN_PAYLOAD)
    login_resp.raise_for_status()

    export_resp = session.post(f"{BASE_URL}/process.php", data=EXPORT_PAYLOAD)
    export_resp.raise_for_status()

    if "zip" not in export_resp.headers.get("Content-Type", "") and not export_resp.content[:2] == b"PK":
        raise RuntimeError("Export failed: response was not a zip file. Check credentials/query parameters.")

    with zipfile.ZipFile(io.BytesIO(export_resp.content)) as zf:
        csv_names = [n for n in zf.namelist() if n.lower().endswith(".csv")]
        if not csv_names:
            raise RuntimeError("No CSV file found inside the exported zip.")
        dest_path = None
        for name in csv_names:
            dest_path = _write_with_fallback_name(zf, name)
            print(f"Saved: {dest_path}")
        return dest_path


def _write_with_fallback_name(zf, member_name):
    base_name = os.path.basename(member_name)
    stem, ext = os.path.splitext(base_name)
    dest_path = os.path.join(ROOT_DIR, base_name)
    counter = 1
    while True:
        try:
            with zf.open(member_name) as src, open(dest_path, "wb") as dest:
                dest.write(src.read())
            return dest_path
        except PermissionError:
            dest_path = os.path.join(ROOT_DIR, f"{stem} ({counter}){ext}")
            counter += 1


def process_csv_to_xlsx(csv_path):
    df = pd.read_csv(csv_path, dtype=str, encoding="utf-8-sig")

    received = pd.to_datetime(df["received_date"], errors="coerce")
    today = datetime.date.today()
    df = df[(received.dt.year == today.year) & (received.dt.month == today.month)]

    df = df.drop(columns=FLEX_FIELDS_TO_DROP)

    if len(df.columns) != len(NEW_HEADERS):
        raise RuntimeError(
            f"Column count mismatch: data has {len(df.columns)} columns after dropping "
            f"flex fields, but {len(NEW_HEADERS)} headers were expected."
        )
    df.columns = NEW_HEADERS

    locations = df["Account Name"].map(normalize_location).tolist()

    stem = "Client Wise All Cases - Bridge (Processed)"
    out_path = os.path.join(ROOT_DIR, f"{stem}.xlsx")
    counter = 1
    while True:
        try:
            df.to_excel(out_path, index=False)
            break
        except PermissionError:
            out_path = os.path.join(ROOT_DIR, f"{stem} ({counter}).xlsx")
            counter += 1
    add_formula_columns(out_path, len(df), locations)

    print(f"Saved: {out_path} ({len(df)} rows)")
    return out_path


def add_formula_columns(out_path, num_rows, locations):
    """Append Hours / New Hours - Days / New Bucket / Location / Inflow Time columns.

    Column letters in the formulas (AN, AC, AO, AF) refer to the fixed source
    layout: AN = DQC_RELEASED_DATE, AC = "I verify date & time", AF =
    API_Created_stamp, and AO is the Hours column itself (added as the 41st
    column, right after AN). Location (AR) is a static value derived from
    Account Name (column R) via LOCATION_MAP, since some source values are
    typos/abbreviations that a lookup formula can't cleanly express.
    """
    wb = load_workbook(out_path)
    ws = wb.active

    hours_col, days_col, bucket_col, location_col, inflow_col = 41, 42, 43, 44, 45  # AO, AP, AQ, AR, AS
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(fill_type="solid", start_color="1F6FC5", end_color="1F6FC5")
    center = Alignment(horizontal="center", vertical="center")
    header_border = Border(bottom=Side(style="medium", color="0D3B66"))

    headers = (
        (hours_col, "Hours"),
        (days_col, "New Hours - Days"),
        (bucket_col, "New Bucket"),
        (location_col, "Location"),
        (inflow_col, "Inflow Time"),
    )
    for col, title in headers:
        ws.cell(row=1, column=col, value=title)

    for col in range(1, inflow_col + 1):
        header_cell = ws.cell(row=1, column=col)
        header_cell.font = header_font
        header_cell.fill = header_fill
        header_cell.alignment = center
        header_cell.border = header_border

    for r in range(2, num_rows + 2):
        hours_cell = ws.cell(
            row=r, column=hours_col,
            value=f'=IF(AN{r}="","DQC Date time Blank",IF(AC{r}="NA","I Verify Date and time Blank",IF(AN{r}="","",(AN{r}-AC{r}))))',
        )
        hours_cell.number_format = "[h]:mm:ss"
        hours_cell.alignment = center

        days_cell = ws.cell(
            row=r, column=days_col,
            value=f'=IF(AN{r}="","",INT(AN{r}-AC{r}))',
        )
        days_cell.alignment = center

        bucket_cell = ws.cell(
            row=r, column=bucket_col,
            value=f'=IF(AN{r}="","DQC Pending",IF(ISNUMBER(AO{r}),IF(AO{r}*24<24,"BELOW 24","ABOVE 24"),AO{r}))',
        )
        bucket_cell.alignment = center

        location_cell = ws.cell(row=r, column=location_col, value=locations[r - 2])
        location_cell.alignment = center

        inflow_cell = ws.cell(
            row=r, column=inflow_col,
            value=(
                f'=IF(OR(AF{r}="",AF{r}="NA"),"Data is not available",'
                f'TEXT(AF{r},"h AM/PM")&" to "&TEXT(AF{r}+TIME(1,0,0),"h AM/PM"))'
            ),
        )
        inflow_cell.alignment = center

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(inflow_col)}{num_rows + 1}"
    ws.row_dimensions[1].height = 22

    fixed_widths = {hours_col: 14, days_col: 16, bucket_col: 16, location_col: 12, inflow_col: 18}
    for col in range(1, inflow_col + 1):
        if col in fixed_widths:
            width = fixed_widths[col]
        else:
            header_text = str(ws.cell(row=1, column=col).value or "")
            width = min(max(len(header_text) + 4, 10), 32)
        ws.column_dimensions[get_column_letter(col)].width = width

    wb.save(out_path)


# Column positions (1-based) in the "Raw data" sheet, used to build the Summary
# Sheet pivots below. Referenced by position (not name) because "Value" occurs
# twice in NEW_HEADERS (columns M and O) and Excel PivotFields can only be
# addressed unambiguously by position when source column names collide.
COL_PROCESS_NAME = 3
COL_CASE_ARS_NO = 5
COL_RECEIVED_DATE = 6
COL_EMPLOYEE_TYPE = 9
COL_VALUE_ENTITY = 15  # second "Value" column - holds entity text (e.g. "Wipro Technologies")
COL_DATA_SOURCE = 38
COL_NEW_BUCKET = 43
COL_LOCATION = 44
COL_INFLOW_TIME = 45

# Each entry: (anchor column, title, row fields, single-item filter or None).
# A filter of (field_col, item_value) isolates that one item in the field's
# filter dropdown (all other items unchecked), matching the reference file.
PIVOT_SPECS = [
    ("A", "Over All Month Inflow Date wise", [COL_RECEIVED_DATE], None),
    ("D", "Over All Month Inflow Date wise", [COL_RECEIVED_DATE], None),
    ("G", "Over All Month Inflow Time wise", [COL_INFLOW_TIME], None),
    ("J", "Over All Entity wise data", [COL_VALUE_ENTITY, COL_EMPLOYEE_TYPE], None),
    ("M", "Over All Entity wise data", [COL_PROCESS_NAME, COL_RECEIVED_DATE], (COL_PROCESS_NAME, "Digital")),
    ("P", "Wipro Cases Count with Location", [COL_LOCATION], None),
    ("S", "Overall Source", [COL_DATA_SOURCE], None),
    ("V", "QC Released Time", [COL_NEW_BUCKET], None),
    ("Y", "QC Released Time", [COL_PROCESS_NAME, COL_DATA_SOURCE, COL_RECEIVED_DATE], (COL_PROCESS_NAME, "GOOGLE")),
]


def _isolate_pivot_item(pivot_field, item_value):
    """Show only the pivot item matching item_value (case-insensitive),
    hiding every other item in that field - mirrors Excel's "filter to one
    item" via the field's dropdown.
    """
    target = None
    for item in pivot_field.PivotItems():
        if str(item.Caption).strip().upper() == item_value.strip().upper():
            target = item
            break
    if target is None:
        raise RuntimeError(f'Pivot item "{item_value}" not found in field "{pivot_field.Name}".')

    target.Visible = True
    for item in pivot_field.PivotItems():
        if item.Name != target.Name:
            item.Visible = False


def _rgb_to_com_color(hex_color):
    """Convert an 'RRGGBB' hex string to the BGR-packed long Excel COM expects."""
    r, g, b = (int(hex_color[i:i + 2], 16) for i in (0, 2, 4))
    return b * 65536 + g * 256 + r


_TITLE_BLUE = _rgb_to_com_color("1F6FC5")


def _style_pivot_table(summary_sheet, anchor_col, title, pivot_table):
    """Give one pivot table a title banner, a built-in pivot style, and a
    border around its body - matches the blue header theme used on the Raw
    data sheet's Hours/Location/Inflow Time columns.
    """
    end_col = chr(ord(anchor_col) + 1)
    title_range = summary_sheet.Range(f"{anchor_col}1:{end_col}1")
    title_range.Merge()
    title_range.Value = title
    title_range.Font.Bold = True
    title_range.Font.Color = 0xFFFFFF  # white (R=G=B, order-independent)
    title_range.Interior.Color = _TITLE_BLUE
    title_range.HorizontalAlignment = -4108  # xlCenter
    title_range.VerticalAlignment = -4108  # xlCenter
    title_range.RowHeight = 20

    try:
        pivot_table.TableStyle2 = "PivotStyleMedium2"
    except Exception:
        pass

    try:
        title_range.BorderAround(LineStyle=1, Weight=2)
        pivot_table.TableRange2.BorderAround(LineStyle=1, Weight=2)
    except Exception:
        pass


def build_summary_pivots(xlsx_path):
    """Rebuild the "Summary Sheet" pivot tables, modeled on the equivalent
    tables in "Wipro Limited - Dashboard.xlsx" (used only as a layout
    reference - this workbook has no link/connection to that file).

    Drives Excel via COM (pywin32) since openpyxl cannot create native
    PivotTable objects. Attaches to the workbook if it's already open in a
    running Excel instance (leaving that instance as-is on exit); otherwise
    opens it in a temporary invisible Excel instance and closes it after
    saving.

    Calls CoInitialize/CoUninitialize around the COM session, since the
    calling thread (e.g. a Streamlit ScriptRunner thread) may not have COM
    initialized on it the way the main thread of a plain script does.
    """
    import pythoncom
    import win32com.client

    xlDatabase = 1
    xlRowField = 1
    xlCount = -4112

    abs_path = os.path.abspath(xlsx_path)
    excel = None
    created_excel = False
    opened_workbook = False

    pythoncom.CoInitialize()
    try:
        try:
            excel = win32com.client.GetActiveObject("Excel.Application")
        except Exception:
            excel = win32com.client.Dispatch("Excel.Application")
            excel.Visible = False
            created_excel = True

        wb = None
        for open_wb in excel.Workbooks:
            if os.path.abspath(open_wb.FullName) == abs_path:
                wb = open_wb
                break
        if wb is None:
            wb = excel.Workbooks.Open(abs_path)
            opened_workbook = True

        prev_alerts = excel.DisplayAlerts
        excel.DisplayAlerts = False
        try:
            raw_sheet = None
            for sheet in wb.Sheets:
                if sheet.Name == "Raw data":
                    raw_sheet = sheet
                    break
            if raw_sheet is None:
                raw_sheet = wb.Sheets(1)
                raw_sheet.Name = "Raw data"

            for sheet in wb.Sheets:
                if sheet.Name == "Summary Sheet":
                    sheet.Delete()

            last_row = raw_sheet.UsedRange.Rows.Count
            source_range = f"'Raw data'!A1:AS{last_row}"

            summary_sheet = wb.Sheets.Add(Before=raw_sheet)
            summary_sheet.Name = "Summary Sheet"

            for i, (anchor_col, title, row_fields, filter_spec) in enumerate(PIVOT_SPECS):
                pivot_cache = wb.PivotCaches().Create(SourceType=xlDatabase, SourceData=source_range)
                pivot_table = pivot_cache.CreatePivotTable(
                    TableDestination=summary_sheet.Range(f"{anchor_col}2"),
                    TableName=f"SummaryPivot{i + 1}",
                )
                for field_col in row_fields:
                    pivot_table.PivotFields(field_col).Orientation = xlRowField
                pivot_table.AddDataField(pivot_table.PivotFields(COL_CASE_ARS_NO), "Count of case_ars_no", xlCount)

                if filter_spec:
                    field_col, item_value = filter_spec
                    _isolate_pivot_item(pivot_table.PivotFields(field_col), item_value)

                _style_pivot_table(summary_sheet, anchor_col, title, pivot_table)

            summary_sheet.Columns.AutoFit()
            try:
                summary_sheet.Activate()
                excel.ActiveWindow.DisplayGridlines = False
            except Exception:
                pass

            wb.Save()
        finally:
            excel.DisplayAlerts = prev_alerts
            if opened_workbook:
                wb.Close(SaveChanges=True)
            if created_excel:
                excel.Quit()
    finally:
        pythoncom.CoUninitialize()


if __name__ == "__main__":
    csv_path = export_csv()
    out_path = process_csv_to_xlsx(csv_path)
    build_summary_pivots(out_path)
